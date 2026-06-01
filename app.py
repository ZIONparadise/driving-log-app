import re
from io import BytesIO

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.formatting.rule import FormulaRule


TARGET_SHEET_KEYWORD = "운행기록부"
START_ROW = 15
END_ROW = 61

RED_FONT = Font(color="FF0000")


def to_int(value):
    if pd.isna(value) or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).replace(",", "").strip()
    if not text:
        return None
    return int(float(text))


def set_cell_value(ws, cell_address, value, font=None):
    target = ws[cell_address]

    for merged_range in ws.merged_cells.ranges:
        if target.coordinate in merged_range:
            cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
            cell.value = value
            if font is not None:
                cell.font = font
            return cell

    target.value = value
    if font is not None:
        target.font = font
    return target


def apply_red_font(ws, cell_address):
    target = ws[cell_address]

    for merged_range in ws.merged_cells.ranges:
        if target.coordinate in merged_range:
            cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
            cell.font = RED_FONT
            return

    target.font = RED_FONT


def read_source_excel(source_file):
    source_file.seek(0)
    filename = source_file.name.lower()

    if filename.endswith(".xls"):
        try:
            source_file.seek(0)
            return pd.read_excel(
                source_file,
                header=None,
                dtype=object,
                engine="xlrd",
                engine_kwargs={"ignore_workbook_corruption": True},
            )
        except Exception:
            source_file.seek(0)
            raw = source_file.read()
            try:
                html_text = raw.decode("utf-8")
            except UnicodeDecodeError:
                html_text = raw.decode("cp949", errors="ignore")
            tables = pd.read_html(html_text)
            if not tables:
                raise ValueError("운행내역 표를 찾을 수 없습니다.")
            return tables[0]

    source_file.seek(0)
    return pd.read_excel(source_file, header=None, dtype=object, engine="openpyxl")


def find_period(df):
    text = " ".join(str(x) for x in df.values.flatten() if not pd.isna(x))
    m = re.search(r"(\d{4}[./-]\d{2}[./-]\d{2})\s*~\s*(\d{4}[./-]\d{2}[./-]\d{2})", text)
    if not m:
        return None, None
    start = m.group(1).replace("-", ".").replace("/", ".")
    end = m.group(2).replace("-", ".").replace("/", ".")
    return start, end


def get_korean_month(period_start, source_filename=""):
    if period_start:
        m = re.search(r"\d{4}[./-](\d{1,2})[./-]\d{1,2}", period_start)
        if m:
            return f"{int(m.group(1))}월"

    m = re.search(r"\d{4}[-./](\d{1,2})[-./]\d{1,2}", source_filename)
    if m:
        return f"{int(m.group(1))}월"

    return "업데이트"


def make_download_filename(period_start, source_filename=""):
    month_label = get_korean_month(period_start, source_filename)
    return f"업무용운행기록부_30th_홍성학_{month_label}.xlsx"


def calculate_distance_values(start_km, end_km, commute):
    distance = None
    business = None

    if start_km is not None and end_km is not None:
        distance = end_km - start_km

    if commute is None:
        commute = 0

    if distance is not None:
        business = distance - commute

    return distance, commute, business


def extract_records(source_file):
    df = read_source_excel(source_file)
    period_start, period_end = find_period(df)

    records = []
    date_pattern = re.compile(r"^\d{4}-\d{2}-\d{2}\([월화수목금토일]\)")

    for _, row in df.iterrows():
        first = row.iloc[0]
        if pd.isna(first):
            continue

        date_text = str(first).strip()
        if not date_pattern.match(date_text):
            continue

        weekday_match = re.search(r"\(([월화수목금토일])\)", date_text)
        weekday = weekday_match.group(1) if weekday_match else ""

        dept = row.iloc[1] if len(row) > 1 and not pd.isna(row.iloc[1]) else ""
        name = row.iloc[2] if len(row) > 2 and not pd.isna(row.iloc[2]) else ""

        start_km = to_int(row.iloc[3]) if len(row) > 3 else None
        end_km = to_int(row.iloc[4]) if len(row) > 4 else None
        commute = to_int(row.iloc[6]) if len(row) > 6 else 0

        distance, commute, business = calculate_distance_values(start_km, end_km, commute)

        purpose = row.iloc[8] if len(row) > 8 and not pd.isna(row.iloc[8]) else ""
        note = row.iloc[9] if len(row) > 9 and not pd.isna(row.iloc[9]) else ""
        final_note = note if str(note).strip() else purpose

        records.append({
            "date": date_text,
            "weekday": weekday,
            "dept": dept,
            "name": name,
            "start_km": start_km,
            "end_km": end_km,
            "distance": distance,
            "commute": commute,
            "business": business,
            "note": final_note,
        })

    return period_start, period_end, records


def get_target_sheet(wb):
    for ws in wb.worksheets:
        if TARGET_SHEET_KEYWORD in ws.title:
            return ws
    return wb.active


def add_conditional_formatting(ws):
    red_font_rule = Font(color="FF0000", bold=True)

    for row_num in range(START_ROW, END_ROW + 1):
        rule_distance_order = FormulaRule(
            formula=[f'AND(ISNUMBER($J{row_num}),ISNUMBER($N{row_num}),$N{row_num}<$J{row_num})'],
            font=red_font_rule,
        )
        ws.conditional_formatting.add(f"J{row_num}", rule_distance_order)
        ws.conditional_formatting.add(f"N{row_num}", rule_distance_order)
        ws.conditional_formatting.add(f"R{row_num}", rule_distance_order)

        rule_commute_over = FormulaRule(
            formula=[f'AND(ISNUMBER($R{row_num}),ISNUMBER($W{row_num}),$W{row_num}>$R{row_num})'],
            font=red_font_rule,
        )
        ws.conditional_formatting.add(f"R{row_num}", rule_commute_over)
        ws.conditional_formatting.add(f"W{row_num}", rule_commute_over)
        ws.conditional_formatting.add(f"AA{row_num}", rule_commute_over)

        rule_usage_sum = FormulaRule(
            formula=[f'AND(ISNUMBER($R{row_num}),ISNUMBER($W{row_num}),ISNUMBER($AA{row_num}),$W{row_num}+$AA{row_num}<>$R{row_num})'],
            font=red_font_rule,
        )
        ws.conditional_formatting.add(f"R{row_num}", rule_usage_sum)
        ws.conditional_formatting.add(f"W{row_num}", rule_usage_sum)
        ws.conditional_formatting.add(f"AA{row_num}", rule_usage_sum)


def enable_excel_recalculation(wb):
    """
    openpyxl은 수식 결과값을 직접 계산하지 않습니다.
    따라서 Excel에서 파일을 열 때 자동 계산되도록 설정합니다.
    """
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass


def update_template(template_file, source_file):
    period_start, period_end, records = extract_records(source_file)

    template_file.seek(0)
    wb = load_workbook(template_file)
    enable_excel_recalculation(wb)
    ws = get_target_sheet(wb)

    if period_start:
        set_cell_value(ws, "E2", period_start)
    if period_end:
        set_cell_value(ws, "E5", period_end)

    target_cols = ["A", "D", "E", "H", "J", "N", "R", "W", "AA", "AE"]
    for row_num in range(START_ROW, END_ROW + 1):
        for col in target_cols:
            set_cell_value(ws, f"{col}{row_num}", None)

    validation_warnings = []

    for i, rec in enumerate(records):
        row_num = START_ROW + i
        if row_num > END_ROW:
            break

        invalid_distance_order = (
            rec["start_km"] is not None
            and rec["end_km"] is not None
            and rec["end_km"] < rec["start_km"]
        )

        invalid_commute_over_distance = (
            rec["distance"] is not None
            and rec["commute"] is not None
            and rec["commute"] > rec["distance"]
        )

        if invalid_distance_order:
            validation_warnings.append(
                f"{rec['date']}: 주행후 거리가 주행전 거리보다 작습니다."
            )

        if invalid_commute_over_distance:
            validation_warnings.append(
                f"{rec['date']}: 출퇴근용 거리가 전체 주행거리보다 큽니다."
            )

        set_cell_value(ws, f"A{row_num}", rec["date"])
        set_cell_value(ws, f"D{row_num}", rec["weekday"])
        set_cell_value(ws, f"E{row_num}", rec["dept"])
        set_cell_value(ws, f"H{row_num}", rec["name"])
        set_cell_value(ws, f"J{row_num}", rec["start_km"])
        set_cell_value(ws, f"N{row_num}", rec["end_km"])

        # 핵심 수정:
        # 복잡한 IF 수식 대신 단순 수식을 사용합니다.
        # Excel에서 값 수정 시 바로 재계산됩니다.
        set_cell_value(ws, f"R{row_num}", f"=N{row_num}-J{row_num}")
        set_cell_value(ws, f"W{row_num}", rec["commute"])
        set_cell_value(ws, f"AA{row_num}", f"=R{row_num}-W{row_num}")

        set_cell_value(ws, f"AE{row_num}", rec["note"])

        if invalid_distance_order:
            for cell in [f"J{row_num}", f"N{row_num}", f"R{row_num}"]:
                apply_red_font(ws, cell)

        if invalid_commute_over_distance:
            for cell in [f"R{row_num}", f"W{row_num}", f"AA{row_num}"]:
                apply_red_font(ws, cell)

    for row_num in range(START_ROW + len(records), END_ROW + 1):
        set_cell_value(ws, f"J{row_num}", None)
        set_cell_value(ws, f"N{row_num}", None)
        set_cell_value(ws, f"R{row_num}", None)
        set_cell_value(ws, f"W{row_num}", None)
        set_cell_value(ws, f"AA{row_num}", None)

    set_cell_value(ws, "K63", "=SUM(S15:S61)")
    set_cell_value(ws, "W63", "=SUM(W15:W61)")
    set_cell_value(ws, "AE63", "=IFERROR(W63/K63,0)")

    add_conditional_formatting(ws)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    download_filename = make_download_filename(period_start, getattr(source_file, "name", ""))

    return output, records, period_start, period_end, validation_warnings, download_filename


def style_preview(row):
    styles = [""] * len(row)
    invalid_distance_order = (
        pd.notna(row.get("start_km"))
        and pd.notna(row.get("end_km"))
        and row.get("end_km") < row.get("start_km")
    )
    invalid_commute_over_distance = (
        pd.notna(row.get("distance"))
        and pd.notna(row.get("commute"))
        and row.get("commute") > row.get("distance")
    )

    for idx, col in enumerate(row.index):
        if invalid_distance_order and col in ["start_km", "end_km", "distance"]:
            styles[idx] = "color: red; font-weight: bold"
        if invalid_commute_over_distance and col in ["distance", "commute", "business"]:
            styles[idx] = "color: red; font-weight: bold"

    return styles


st.title("업무용승용차 운행기록부 자동 업데이트")

st.write("제출해야 하는 운행기록부 양식과 다운로드 받은 운행내역 파일을 업로드하면 자동으로 내용을 채웁니다.")

st.caption("자동 계산 규칙: 7번 주행거리 = N열 주행후거리 - J열 주행전거리 / 9번 일반 업무용 = R열 주행거리 - W열 출퇴근용")
st.caption("다운로드된 엑셀 파일 안에도 수식이 들어가므로, 주행전/주행후/출퇴근용 값을 수정하면 자동으로 재계산됩니다.")

template_file = st.file_uploader("1. 제출 양식 파일 업로드 (.xlsx)", type=["xlsx"])
source_file = st.file_uploader("2. 다운로드 받은 운행내역 파일 업로드 (.xls 또는 .xlsx)", type=["xls", "xlsx"])

if template_file and source_file:
    if st.button("운행기록부 업데이트"):
        try:
            output, records, period_start, period_end, validation_warnings, download_filename = update_template(template_file, source_file)

            if len(records) == 0:
                st.warning("운행내역을 찾지 못했습니다. 다운로드 파일의 표 구조가 예상과 다를 수 있습니다.")
            else:
                st.success(f"업데이트 완료: {len(records)}건")

            if period_start and period_end:
                st.caption(f"과세기간: {period_start} ~ {period_end}")

            st.caption(f"저장 파일명: {download_filename}")

            if validation_warnings:
                st.warning("검산 확인이 필요한 항목이 있습니다. 다운로드된 엑셀에서도 해당 숫자가 빨간색으로 표시됩니다.")
                for warning in validation_warnings:
                    st.write(f"- {warning}")

            preview_df = pd.DataFrame(records)
            if not preview_df.empty:
                st.dataframe(preview_df.style.apply(style_preview, axis=1), use_container_width=True)
            else:
                st.dataframe(preview_df, use_container_width=True)

            st.download_button(
                label="업데이트된 엑셀 다운로드",
                data=output,
                file_name=download_filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        except Exception as e:
            st.error("처리 중 오류가 발생했습니다.")
            st.exception(e)
else:
    st.info("먼저 제출 양식과 운행내역 파일을 모두 업로드해 주세요.")
